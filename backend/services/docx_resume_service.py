from io import BytesIO


def extract_docx_text(content):
    from docx import Document

    document = Document(BytesIO(content))
    return "\n".join(
        paragraph.text.strip()
        for paragraph in _all_paragraphs(document)
        if paragraph.text.strip()
    )


def create_tailored_docx(original_path, base_text, approved_text):
    from docx import Document

    document = Document(original_path)
    paragraphs = [paragraph for paragraph in _all_paragraphs(document) if paragraph.text.strip()]
    base_lines = [line.strip() for line in base_text.splitlines() if line.strip()]
    approved_lines = [line.strip() for line in approved_text.splitlines() if line.strip()]

    if base_lines and len(paragraphs) < len(base_lines):
        raise ValueError("The stored DOCX no longer matches this resume version.")

    for index, paragraph in enumerate(paragraphs):
        replacement = approved_lines[index] if index < len(approved_lines) else ""
        _replace_paragraph_text(paragraph, replacement)

    for extra_line in approved_lines[len(paragraphs):]:
        document.add_paragraph(extra_line)

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer


def _all_paragraphs(document):
    for paragraph in document.paragraphs:
        yield paragraph

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph


def _replace_paragraph_text(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.text = text
