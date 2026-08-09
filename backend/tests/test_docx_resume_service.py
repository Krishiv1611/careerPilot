from io import BytesIO
from pathlib import Path
import sys
import tempfile
from unittest import TestCase

from docx import Document


BACKEND_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_ROOT))

from services.docx_resume_service import create_tailored_docx


class DocxResumeServiceTests(TestCase):
    def test_export_keeps_template_structure_while_applying_approved_text(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            original_path = Path(temp_dir) / "resume.docx"
            document = Document()
            document.sections[0].header.paragraphs[0].text = "Student Name"
            heading = document.add_paragraph("Projects", style="Heading 1")
            heading.runs[0].bold = True
            document.add_paragraph("Built a Python API")
            table = document.add_table(rows=1, cols=1)
            table.cell(0, 0).text = "Education"
            document.sections[0].footer.paragraphs[0].text = "Page footer"
            document.save(original_path)

            exported = create_tailored_docx(
                str(original_path),
                "Projects\nBuilt a Python API\nEducation",
                "Projects\nBuilt a FastAPI service\nEducation",
            )
            result = Document(BytesIO(exported.getvalue()))

        self.assertEqual("Student Name", result.sections[0].header.paragraphs[0].text)
        self.assertEqual("Page footer", result.sections[0].footer.paragraphs[0].text)
        self.assertEqual(1, len(result.tables))
        self.assertEqual("Heading 1", result.paragraphs[0].style.name)
        self.assertTrue(result.paragraphs[0].runs[0].bold)
        self.assertEqual("Built a FastAPI service", result.paragraphs[1].text)
        self.assertEqual("Education", result.tables[0].cell(0, 0).text)
